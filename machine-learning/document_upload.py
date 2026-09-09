"""
Upload an institutional document into the retrieval corpus, from the portal.

    POST /document/preview   extract, chunk and report. Writes NOTHING.
    POST /document/apply     embed and store the extraction that was reviewed.

Accepts .md, .txt, .pdf and .docx. The text is chunked with the SAME function
the batch importer uses, so a document ingested here is indistinguishable from
one ingested by document_knowledge_importer.py -- there is one chunking rule,
not two that can drift.

WHY THIS EXISTS, AGAINST AN EARLIER DECISION
--------------------------------------------
document_knowledge_importer.py records the opposite choice: "this is a batch
script, not an admin UI. An authenticated file-upload surface is unmentioned in
the thesis, adds attack surface, and competes for time." That was right while
the corpus was two files that changed once. It stopped being right once the
registrar began issuing a revised calendar and handbook every year: a corpus
that can only be updated by someone with a terminal is a corpus that stops being
updated.

The attack surface is answered rather than ignored: the endpoint is behind the
admin/researcher role, the extraction is shown before anything is written, and
the file itself is never stored or served -- only text that a person approved.

PROVENANCE IS NOT OPTIONAL
--------------------------
Every document row carries source_origin, provided_by and a SHA-256 of the
uploaded bytes. "Where did this claim come from" has to be answerable from the
database alone, without the original file, or the corpus is unauditable.
"""

from __future__ import annotations

import base64
import binascii
import hashlib
import io
import re
import sys

from flask import Blueprint, jsonify, request

import database_connector as db
import document_knowledge_importer as imp

document_bp = Blueprint("document", __name__)


def _use_in_process_model() -> None:
    """
    Chunk and embed with the model already resident in this process.

    document_knowledge_importer talks to the ML service over HTTP, which is
    right for a batch script running outside it. Inside the service that means
    the process issuing HTTP requests to itself: it burns one of four waitress
    threads per call and turns a local function call into a network round trip
    that can time out against its own backlog.

    The import is deferred because ai_api_service imports THIS module while
    building its blueprint list, and the embedder is not loaded at that point
    anyway. If anything fails the HTTP path is left in place -- slower, but
    still correct.
    """
    try:
        # NOT `import ai_api_service`. The service is started as a script, so it
        # is registered under '__main__'; importing it by name loads a SECOND
        # copy of the module, which reloads the sentence-transformer and
        # re-registers the blueprints on a fresh Flask app. Look the running
        # module up instead.
        svc = sys.modules.get("ai_api_service") or sys.modules.get("__main__")
        if svc is None or getattr(svc, "_embedder", None) is None:
            return
        tok = svc._embedder.tokenizer
        imp.count_tokens = lambda texts: [
            len(tok.encode(str(t), add_special_tokens=True)) for t in texts]
        imp.embed_batch = lambda texts: [
            v.tolist() for v in svc._embed(list(texts))]
    except Exception:
        pass

MAX_UPLOAD_BYTES = 25 * 1024 * 1024
SUPPORTED = (".md", ".txt", ".pdf", ".docx")


# ------------------------------------------------------------- extraction
def _from_pdf(raw: bytes) -> str:
    import fitz                                   # PyMuPDF
    with fitz.open(stream=raw, filetype="pdf") as doc:
        return "\n\n".join(page.get_text("text") for page in doc)


def _from_docx(raw: bytes) -> str:
    import docx
    document = docx.Document(io.BytesIO(raw))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:                 # calendars often live in tables
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_text(filename: str, raw: bytes) -> str:
    suffix = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix not in SUPPORTED:
        raise ValueError(
            f"unsupported file type '{suffix or filename}'. "
            f"Accepted: {', '.join(SUPPORTED)}")

    if suffix == ".pdf":
        text = _from_pdf(raw)
    elif suffix == ".docx":
        text = _from_docx(raw)
    else:
        text = raw.decode("utf-8", errors="replace")

    # Collapse the ragged whitespace a PDF extraction leaves behind, without
    # destroying paragraph breaks -- the chunker splits on those.
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _read(payload) -> tuple[str, bytes]:
    filename = str(payload.get("filename") or "").strip()
    if not filename:
        raise ValueError("`filename` is required to determine the file type")
    raw_b64 = payload.get("content_b64")
    if not isinstance(raw_b64, str) or not raw_b64.strip():
        raise ValueError("`content_b64` is required")
    try:
        raw = base64.b64decode(raw_b64, validate=True)
    except (binascii.Error, ValueError):
        raise ValueError("`content_b64` is not valid base64")
    if len(raw) > MAX_UPLOAD_BYTES:
        raise ValueError(f"file exceeds {MAX_UPLOAD_BYTES // (1024*1024)} MB")
    return filename, raw


def _title_of(filename: str, override: str | None) -> str:
    if override and override.strip():
        return override.strip()[:200]
    return filename.rsplit(".", 1)[0].strip()[:200]


def _summarise(filename, raw, text, title, doc_type) -> dict:
    chunks = imp.chunk_document(text)
    existing = db.fetch_all(
        "select id::text, title, doc_type, source_checksum, ingested_at "
        "from geobot.document where title = %s", (title,))

    return {
        "kind": "document",
        "filename": filename,
        "title": title,
        "doc_type": doc_type,
        "checksum": hashlib.sha256(raw).hexdigest(),
        "characters": len(text),
        "chunks": len(chunks),
        "max_tokens": max((n for _, n in chunks), default=0),
        "ceiling": imp.MAX_TOKENS,
        "replaces": [
            {"title": e["title"], "doc_type": e["doc_type"],
             "ingested_at": e["ingested_at"].isoformat() if e["ingested_at"] else None,
             "same_bytes": e["source_checksum"] == hashlib.sha256(raw).hexdigest()}
            for e in existing
        ],
        "preview": [{"index": i, "tokens": n, "text": c[:260]}
                    for i, (c, n) in enumerate(chunks[:6])],
    }


# -------------------------------------------------------------- endpoints
@document_bp.post("/document/preview")
def preview():
    _use_in_process_model()
    payload = request.get_json(silent=True) or {}
    try:
        filename, raw = _read(payload)
        text = extract_text(filename, raw)
    except ValueError as exc:
        return jsonify(error=str(exc)), 400
    except Exception as exc:
        return jsonify(error=f"could not read the file: {exc}"), 400

    if len(text) < 40:
        return jsonify(error="almost no text could be extracted. A scanned PDF with "
                             "no text layer has to be put through OCR first."), 400

    title = _title_of(filename, payload.get("title"))
    doc_type = payload.get("doc_type") or imp.infer_doc_type(filename)
    return jsonify(**_summarise(filename, raw, text, title, doc_type))


@document_bp.post("/document/apply")
def apply():
    _use_in_process_model()
    payload = request.get_json(silent=True) or {}
    expected = str(payload.get("checksum") or "").strip()
    if not expected:
        return jsonify(error="`checksum` from /document/preview is required"), 400
    try:
        filename, raw = _read(payload)
        text = extract_text(filename, raw)
    except ValueError as exc:
        return jsonify(error=str(exc)), 400
    except Exception as exc:
        return jsonify(error=f"could not read the file: {exc}"), 400

    if hashlib.sha256(raw).hexdigest() != expected:
        return jsonify(error="checksum mismatch: this file does not match the one "
                             "that was previewed. Preview it again."), 409

    title = _title_of(filename, payload.get("title"))
    doc_type = payload.get("doc_type") or imp.infer_doc_type(filename)
    checksum = hashlib.sha256(raw).hexdigest()
    summary = _summarise(filename, raw, text, title, doc_type)

    with db.cursor() as cur:
        # Replacing by title keeps one row per document rather than accumulating
        # a new copy on every re-upload -- which is how the calendar ended up in
        # the corpus twice, with retrieval spending its top-k on duplicates.
        cur.execute("""delete from geobot.document_chunk
                        where document_id in (select id from geobot.document
                                               where title = %s)""", (title,))
        cur.execute("delete from geobot.document where title = %s", (title,))
        replaced = cur.rowcount

        cur.execute(
            """insert into geobot.document
                 (title, doc_type, source_origin, provided_by, source_checksum,
                  data_origin)
               values (%s,%s,%s,%s,%s,'real') returning id""",
            (title, doc_type, payload.get("source_origin") or filename,
             payload.get("provided_by") or "admin portal upload", checksum))
        document_id = cur.fetchone()["id"]

    written = imp._write_chunks(document_id, text, "real")

    totals = db.fetch_all(
        "select count(distinct d.id) docs, count(c.id) chunks "
        "from geobot.document d left join geobot.document_chunk c "
        "on c.document_id = d.id")[0]

    return jsonify(applied=True, kind="document", title=title, doc_type=doc_type,
                   chunks_written=written, replaced_documents=replaced,
                   corpus={"documents": totals["docs"], "chunks": totals["chunks"]},
                   summary=summary)
